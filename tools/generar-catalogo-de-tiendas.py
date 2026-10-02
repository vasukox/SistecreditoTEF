# -*- coding: utf-8 -*-
"""
GENERA EL CATALOGO DE TIENDAS QUE VIAJA DENTRO DEL APK.

    python tools/generar-catalogo-de-tiendas.py [ruta-al-xlsx]

Lee la hoja de StoreId de Sistecredito y reescribe

    src/SistecreditoTEF.Maui/Services/Tiendas/CatalogoDeTiendas.Datos.cs

Despues hay que CORRER LAS PRUEBAS y RECOMPILAR el APK: el catalogo va embebido,
no se descarga.

    dotnet test tests/SistecreditoTEF.Tests/SistecreditoTEF.Tests.csproj
    dotnet build src/SistecreditoTEF.Maui/SistecreditoTEF.Maui.csproj -c Release -f net10.0-android -p:Ambiente=produccion

───────────────────────────────────────────────────────────────────────────────
POR QUE ESTE SCRIPT ESTA EN EL REPO Y NO EN LA CARPETA DE ALGUIEN
───────────────────────────────────────────────────────────────────────────────
El StoreId decide a nombre de que tienda queda cada credito en Sistecredito.
Creditos hechos en una tienda de Suba aparecieron registrados como tienda 037
porque ese valor estaba horneado en el APK. La leccion no es solo "no hornear un
StoreId": es que el camino desde la hoja oficial hasta el APK tiene que ser
REPETIBLE por cualquiera, y no depender de que alguien se acuerde de como lo hizo.

Si el dia de mañana llegan 200 tiendas mas en otro archivo, esto es lo unico que
hay que correr.

───────────────────────────────────────────────────────────────────────────────
QUE VALIDA ANTES DE ESCRIBIR
───────────────────────────────────────────────────────────────────────────────
Falla —y no escribe nada— si:

  · el codigo de tienda no son 3 digitos,
  · el StoreId no es un ObjectId de 24 caracteres hexadecimales,
  · hay codigos o StoreId repetidos.

Un StoreId mal copiado no da error en ningun lado: simplemente manda los creditos
de una tienda a otra. Por eso se revienta aca, en la maquina de quien genera, y no
en una caja.

El NOMBRE si puede faltar: varias tiendas de la hoja llegan sin el. Se conservan
igual —tienen StoreId valido— y la app las muestra por codigo. Sacarlas obligaria
a teclear 24 caracteres a mano justo en las tiendas que se estan abriendo.

Requiere: pip install openpyxl
"""
import os
import re
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl. Instalalo con:  python -m pip install openpyxl")

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

XLSX_POR_DEFECTO = os.path.join(RAIZ, "CREDENCIALES Y PARAMETROS TEF-PERMODA (1).xlsx")
DEST = os.path.join(
    RAIZ, "src", "SistecreditoTEF.Maui", "Services", "Tiendas", "CatalogoDeTiendas.Datos.cs")

CODIGO = re.compile(r"^\d{3}$")
OBJECT_ID = re.compile(r"^[0-9a-f]{24}$")


def elegir_hoja(wb):
    """La hoja de StoreId, buscada por nombre y no por posicion.

    Las hojas de un .xlsx se indexan por rId, no por numero, asi que tomar "la
    sexta" es una forma comoda de leer la hoja equivocada. Se busca por nombre.
    """
    for nombre in wb.sheetnames:
        plano = nombre.upper().replace(" ", "").replace("-", "").replace("_", "")
        if "STOREID" in plano and "SISTECREDITO" in plano:
            return wb[nombre]

    sys.exit(
        "No se encontro una hoja de StoreId de Sistecredito.\n"
        "Hojas del archivo: " + ", ".join(wb.sheetnames))


def ubicar_columnas(ws):
    """Encuentra las columnas por ENCABEZADO, no por posicion.

    Si el proximo archivo trae las columnas en otro orden —o una columna nueva en
    el medio— leerlas por indice fijo produciria un catalogo silenciosamente mal
    armado. Buscarlas por nombre falla ruidosamente si el formato cambia de verdad.
    """
    encabezados = [str(c).strip().upper() if c is not None else "" for c in next(
        ws.iter_rows(min_row=1, max_row=1, values_only=True))]

    def buscar(*candidatos):
        for i, h in enumerate(encabezados):
            plano = h.replace(" ", "")
            if any(c in plano for c in candidatos):
                return i
        return None

    i_codigo = buscar("TIENDA", "CODIGO", "CÓDIGO")
    i_store = buscar("STOREID")
    i_ciudad = buscar("CIUDAD")

    if i_codigo is None or i_store is None:
        sys.exit(
            "No se reconocieron las columnas de la hoja.\n"
            f"Encabezados encontrados: {encabezados}\n"
            "Se esperaba una columna de tienda/codigo y una de Store ID.")

    # La segunda columna "TIENDA" suele ser el nombre. Si existe, se usa.
    i_nombre = None
    for i, h in enumerate(encabezados):
        if i != i_codigo and "TIENDA" in h.replace(" ", "").upper():
            i_nombre = i
            break
    if i_nombre is None:
        i_nombre = buscar("NOMBRE")

    return i_codigo, i_nombre, i_ciudad, i_store


def limpiar_nombre(valor, codigo):
    """Nombre de la tienda, sin basura de referencias de celda y sin el codigo repetido."""
    if valor is None:
        return ""
    n = re.sub(r"[+$][A-Z0-9:$]+", "", str(valor)).strip()
    n = re.sub(r"\s+", " ", n)
    if n.startswith(codigo + " "):
        n = n[len(codigo) + 1:]
    return n


def leer(ruta):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    ws = elegir_hoja(wb)
    i_cod, i_nom, i_ciu, i_sid = ubicar_columnas(ws)

    filas, errores = [], []
    for n, fila in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if fila is None or all(c is None for c in fila):
            continue

        cod = "" if fila[i_cod] is None else str(fila[i_cod]).strip()
        sid = "" if fila[i_sid] is None else str(fila[i_sid]).strip()
        ciu = "" if i_ciu is None or fila[i_ciu] is None else str(fila[i_ciu]).strip()
        nom = limpiar_nombre(fila[i_nom] if i_nom is not None else None, cod)

        if not CODIGO.match(cod):
            errores.append(f"  fila {n}: codigo de tienda invalido: {cod!r}")
            continue
        if not OBJECT_ID.match(sid):
            errores.append(f"  fila {n} (tienda {cod}): StoreId invalido: {sid!r}")
            continue

        filas.append((cod, nom, ciu or "(sin ciudad)", sid))

    repetidos = [c for c in {f[0] for f in filas} if [x[0] for x in filas].count(c) > 1]
    if repetidos:
        errores.append(f"  codigos de tienda repetidos: {sorted(repetidos)}")

    ids = [f[3] for f in filas]
    duplicados = sorted({i for i in ids if ids.count(i) > 1})
    if duplicados:
        errores.append(f"  StoreId repetidos: {duplicados}")

    if errores:
        sys.exit("La hoja tiene problemas y NO se genero nada:\n" + "\n".join(errores))

    if not filas:
        sys.exit("La hoja no tiene ninguna tienda utilizable.")

    filas.sort(key=lambda f: f[0])
    return filas


def escribir(filas, origen):
    def esc(s):
        return s.replace("\\", "\\\\").replace('"', '\\"')

    a_cod = max(len(f[0]) for f in filas)
    a_nom = max(len(f[1]) for f in filas)
    a_ciu = max(len(f[2]) for f in filas)

    lineas = []
    for cod, nom, ciu, sid in filas:
        c = ('"%s",' % esc(cod)).ljust(a_cod + 4)
        n = ('"%s",' % esc(nom)).ljust(a_nom + 4)
        u = ('"%s",' % esc(ciu)).ljust(a_ciu + 4)
        lineas.append('        new(%s %s %s "%s"),' % (c, n, u, sid))

    sin_nombre = sum(1 for f in filas if not f[1])

    contenido = """// <auto-generated>
//     ARCHIVO GENERADO. NO EDITAR A MANO.
//
//     Origen : %s
//     Filas  : %d  (%d sin nombre en la hoja)
//     Script : tools/generar-catalogo-de-tiendas.py
//
//     Para regenerarlo con una hoja nueva:
//         python tools/generar-catalogo-de-tiendas.py <ruta-al-xlsx>
//     y despues correr las pruebas y RECOMPILAR el APK: esto va embebido.
//
//     La logica vive en CatalogoDeTiendas.cs; este archivo solo trae los datos,
//     justamente para que regenerar no pise codigo.
// </auto-generated>

namespace SistecreditoTEF.Maui.Services.Tiendas;

public static partial class CatalogoDeTiendas
{
    /// <summary>Las tiendas de la hoja oficial, ordenadas por codigo.</summary>
    private static readonly TiendaDelCatalogo[] Datos =
    [
%s
    ];
}
""" % (os.path.basename(origen), len(filas), sin_nombre, "\n".join(lineas))

    os.makedirs(os.path.dirname(DEST), exist_ok=True)
    with open(DEST, "w", encoding="utf-8") as f:
        f.write(contenido)

    return sin_nombre


def main():
    origen = sys.argv[1] if len(sys.argv) > 1 else XLSX_POR_DEFECTO
    if not os.path.exists(origen):
        sys.exit(f"No existe el archivo: {origen}")

    filas = leer(origen)
    sin_nombre = escribir(filas, origen)

    print(f"Origen  : {origen}")
    print(f"Escrito : {DEST}")
    print(f"Tiendas : {len(filas)}  (sin nombre en la hoja: {sin_nombre})")
    print(f"Primera : {filas[0][0]} {filas[0][1]}")
    print(f"Ultima  : {filas[-1][0]} {filas[-1][1]}")
    print()
    print("Falta: correr las pruebas y recompilar el APK.")


if __name__ == "__main__":
    main()
